#!/usr/bin/env python3
"""Create the editable human review document from the actual display decisions."""
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'data/eastern-jin-display/machine-decisions.json'
OUTPUT = ROOT / 'outputs/东晋待考项目机器判断_20261011.docx'
LABELS = {'adopted': '暂按此项采用', 'partial': '部分采用', 'suspended': '暂缓采用',
          'retained_with_note': '保留原列并注明'}
KINDS = {'date': '年代', 'parentage': '行政归属', 'identity': '实体身份',
         'text': '文本与校字', 'other': '其他问题', 'source_access': '版本与来源'}


def unique(values):
    return list(dict.fromkeys(x for x in values if x))


def span(start, end):
    a = '未详' if start is None else str(start)
    b = '未详' if end is None else str(end)
    return a if a == b else f'{a}—{b}'


def years_text(years):
    yrs = sorted(set(years or []))
    if not yrs:
        return '未定年'
    ranges = []
    first = last = yrs[0]
    for y in yrs[1:]:
        if y == last + 1:
            last = y
        else:
            ranges.append(span(first, last))
            first = last = y
    ranges.append(span(first, last))
    return '、'.join(ranges)


def add_field(paragraph, text):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), text)
    paragraph._p.append(field)


def add_line(doc, label, text, small=False):
    if not text:
        return
    p = doc.add_paragraph(style='证据' if small else '正文条目')
    if label == '人工意见':
        p.paragraph_format.keep_with_next = True
    p.add_run(label + '  ').bold = True
    p.add_run(clean_text(text))
    return p


def clean_text(value):
    text = str(value)
    text = text.replace('\\n', '\n')
    text = re.sub(r'\\underset\s*\{\\cdot\}\s*\{([^{}]*)\}', r'\1', text)
    text = re.sub(r'\\(?:underline|text|mathrm|mathbf)\s*\{', '', text)
    text = re.sub(r'[\^_]\s*\{([^{}]*)\}', r'\1', text)
    text = text.replace('$', '').replace('{', '').replace('}', '')
    text = re.sub(r'#+\s*', '', text)
    text = re.sub(r'<[^>]+>', ' ', text)
    # Some Work excerpts end midway through an OCR table tag.
    text = re.sub(r'<(?:/?[A-Za-z][^<>]*?)?(?=……|$)', '', text)
    return re.sub(r'[ \t]+', ' ', text).strip()


def sources_text(record):
    out = []
    for e in record.get('evidence', []):
        name = e.get('source') or e.get('source_title') or '来源未题'
        if isinstance(name, dict):
            name = name.get('title') or name.get('id') or str(name)
        loc = e.get('locator') or {}
        chapter = e.get('chapter') or loc.get('chapter') or ''
        page = e.get('book_page') or loc.get('book_page')
        ref = ' '.join(str(x) for x in [name, chapter, f'书页{page}' if page else None] if x)
        quote = clean_text(e.get('excerpt') or e.get('raw_excerpt') or '').replace('\n', ' ').strip()
        # Long extracts stay in the companion machine ledger; the document prints a located excerpt.
        if len(quote) > 125:
            quote = quote[:125] + '……'
        out.append(f'{ref}：{quote}' if quote else ref)
    return unique(out)


def candidates_text(record):
    output = []
    for c in record.get('candidates', []):
        if not isinstance(c, dict):
            output.append(str(c))
            continue
        phases = []
        for p in c.get('phases', []):
            raw = p.get('raw') or p.get('period_raw') or span(p.get('start'), p.get('end'))
            starts = p.get('start_candidates') or []
            ends = p.get('end_candidates') or []
            options = []
            if starts:
                options.append('起年候选' + '、'.join(map(str, starts)))
            if ends:
                options.append('终年候选' + '、'.join(map(str, ends)))
            phases.append(str(raw) + ('（' + '；'.join(options) + '）' if options else ''))
        parents = unique(f"{p.get('name') or p.get('parent_name') or p.get('id')} {span(p.get('start'), p.get('end'))}"
                         for p in c.get('parents', []))
        detail = '；'.join(unique(phases))
        if parents:
            detail += ('；' if detail else '') + '上级候选 ' + '、'.join(parents)
        if detail:
            output.append((c.get('name') or '') + '：' + detail)
    return '；'.join(output)


def adoption_text(record):
    details = unique(f"{p.get('name') or p.get('entity_name') or ''} {span(p.get('start'), p.get('end'))}"
                     for p in record.get('adopted_periods', []))
    details += unique(f"{p.get('entity_name') or ''} → {p.get('parent_name') or p.get('parent_id') or '归属待考'} {span(p.get('start'), p.get('end'))}"
                      for p in record.get('adopted_parentages', []))
    return '；'.join(unique(details))


def main():
    data = json.loads(SOURCE.read_text())
    records = data['decisions']
    structural = data.get('structural_decisions', [])
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.7)
    sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
    sec.header_distance, sec.footer_distance = Cm(.6), Cm(.7)
    styles = doc.styles
    for border in styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for fonts in styles.element.xpath('.//w:rFonts'):
        for attribute in list(fonts.attrib):
            if 'theme' in attribute.lower():
                del fonts.attrib[attribute]
        for attribute in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + attribute), 'Songti SC')
    for key in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        st = styles[key]
        st.font.name = 'Songti SC'
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC')
        st.paragraph_format.widow_control = True
    styles['Normal'].font.size = Pt(10.5)
    styles['Normal'].paragraph_format.line_spacing = 1.13
    styles['Normal'].paragraph_format.space_after = Pt(4)
    styles['Title'].font.size = Pt(23)
    styles['Heading 1'].font.size = Pt(16)
    styles['Heading 2'].font.size = Pt(12)
    styles['Heading 2'].paragraph_format.space_before = Pt(8)
    styles['Heading 2'].paragraph_format.space_after = Pt(4)
    styles['Heading 2'].paragraph_format.keep_with_next = True
    for name, size in [('正文条目', 10.5), ('证据', 9.5)]:
        st = styles.add_style(name, 1)
        st.base_style = styles['Normal']
        st.font.size = Pt(size)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.08
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    add_field(footer, 'PAGE')
    footer.add_run(' 页')
    doc.core_properties.title = '东晋待考项目机器判断与人工核对'
    doc.core_properties.subject = '东晋第一轮年度展示裁定'
    doc.core_properties.author = '六朝年表项目'
    doc.add_paragraph('东晋待考项目机器判断与人工核对', 'Title')
    doc.add_paragraph('317—420年行政区划首版  2026年10月11日')
    counts = Counter(r.get('status') for r in records)
    doc.add_paragraph(f'本稿逐条列出 Work 留下的{len(records)}项开放问题及{len(structural)}组跨年归属问题，供你核对本轮网页采用的年份和州郡县归属。机器判断服务于首版展示，原始研究记录继续保留；采用某项不等于史学问题已经解决。')
    doc.add_paragraph('本轮在来源列出的可用时间中选择最早且不与已知事件冲突的一项；归属优先采用明确行政属隶、同一机构的可靠阶段关系和可支持的候选。没有可靠候选、涉及不同机构身份或仅有寄治原籍的，暂缓作为主表确定关系。侨州、侨郡、侨县的标记独立保留。')
    doc.add_paragraph('判断结果：' + '；'.join(f'{LABELS.get(k,k)} {v}项' for k,v in counts.items()) + '。')
    doc.add_paragraph('“397年后”可自397年展示，并注明非确切始置年；只有418或420年存在观察、没有早年候选的，不一概倒推317年。凡有原始竞争说法，本稿保留问题与候选。引文过长时只列定位节录，完整内容见机器判断JSON及原研究证据。')
    doc.add_paragraph('核对时可直接在每条“人工意见”之后填写修改；引用本稿编号及原问题ID即可回写对应数据。附表的S编号为跨年父级问题，可能与前面的单项问题涉及同一实体，因此不应把两部分相加理解为独立政区数。')
    doc.add_paragraph('逐项判断', 'Heading 1')
    last_kind = None
    for i, r in enumerate(records, 1):
        label = r.get('label') or '、'.join(e.get('name', '') for e in r.get('entities', [])) or '全局资料问题'
        kind = KINDS.get(r.get('kind'), r.get('kind_label') or '其他')
        doc.add_paragraph(f'{i:03d}  {label.replace("（", " ").replace("）", " ")}  {kind}', 'Heading 2')
        add_line(doc, '原问题', r.get('original_question'))
        cand = candidates_text(r)
        if cand:
            add_line(doc, '候选', cand)
        decision = r.get('decision', '')
        adopted = adoption_text(r)
        add_line(doc, '本轮采用', f"{LABELS.get(r.get('status'),r.get('status',''))}。{decision}" + (f' 具体列示：{adopted}' if adopted else ''))
        add_line(doc, '判断理由', r.get('reason'))
        evidence = sources_text(r)
        if evidence:
            add_line(doc, '出处', '；'.join(evidence[:3]), small=True)
            if len(evidence) > 3:
                add_line(doc, '补充出处', f'另有{len(evidence)-3}条证据，按原问题ID见完整机器判断JSON。', small=True)
        add_line(doc, '核对要点', r.get('manual_review'))
        add_line(doc, '人工意见', '________________________________________________')
        add_line(doc, '索引', f"{r.get('issue_id','')}  判断{r.get('id','')}  置信度{r.get('confidence','未定')}", small=True)
    doc.add_page_break()
    doc.add_paragraph('跨年归属问题附表', 'Heading 1')
    doc.add_paragraph('这些问题来自年度树结构检查，包括父级阶段缺口和多候选。主表未采用的归属留待考，不将侨寄地点自动转成行政隶属。')
    for i, r in enumerate(structural, 1):
        doc.add_paragraph(f'S{i:03d}  {r.get("name") or r.get("entity_id","")}', 'Heading 2')
        add_line(doc, '涉及年份', years_text(r.get('years')))
        candidates = r.get('original_candidates', [])
        if candidates:
            add_line(doc, '原候选', '；'.join(str(x.get('name') or x.get('id') or x) if isinstance(x,dict) else str(x) for x in candidates))
        add_line(doc, '本轮采用', f"{LABELS.get(r.get('status'),r.get('status',''))}。{r.get('decision','')}" + ('  ' + adoption_text(r) if adoption_text(r) else ''))
        add_line(doc, '判断理由', r.get('reason'))
        evidence = sources_text(r)
        if evidence:
            add_line(doc, '出处', '；'.join(evidence[:2]), small=True)
        add_line(doc, '核对要点', r.get('manual_review'))
        add_line(doc, '人工意见', '________________________________________________')
        add_line(doc, '索引', f"{r.get('id','')}  {r.get('entity_id','')}  置信度{r.get('confidence','未定')}", small=True)
    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    print(json.dumps({'file': str(OUTPUT), 'items': len(records), 'structural_groups': len(structural), 'status_counts': dict(counts)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
